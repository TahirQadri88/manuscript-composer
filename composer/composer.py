# -*- coding: utf-8 -*-
"""
Manuscript Composer — converts a marked-up transcript (transcript.txt) into a styled .docx
Supports RTL (Arabic, Urdu, Persian, Pashto, Sindhi...) and LTR (English, etc.), footnotes,
uncertain-reading highlights with Word comments, scribal blanks, page markers, poetry verses.

Usage:  python3 composer.py transcript.txt output.docx [template.docx]
Requires: pip install "python-docx>=1.2.0"
"""
import re, sys, copy
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart

# =====================================================================
#  STYLE SHEET  —  EDIT THIS BLOCK TO MATCH THE CLIENT'S STYLE SHEET
# =====================================================================
STYLE = {
    "page": {"width_cm": 21.0, "height_cm": 29.7,          # A4. (Letter = 21.59 x 27.94, A5 = 14.8 x 21)
             "top_cm": 2.5, "bottom_cm": 2.5, "right_cm": 2.5, "left_cm": 2.5,
             "gutter_cm": 0, "mirror_margins": False,
             "header_cm": 1.25, "footer_cm": 1.25},
    "page_number": {"show": True, "position": "footer",      # "footer" | "header"
                    "align": "center", "size": 11},          # "center" | "start" | "end"
    "header_text": "",                                       # running header text (empty = none)
    # Per-language fonts / sizes / spacing.  size in pt, line = multiple (1.0, 1.15, 1.5, 2.0)
    "langs": {
        "ar": {"font": "Traditional Arabic",     "size": 16, "line": 1.2,  "rtl": True,  "tag": "ar-SA"},
        "ur": {"font": "Jameel Noori Nastaleeq", "size": 14, "line": 1.35, "rtl": True,  "tag": "ur-PK"},
        "fa": {"font": "B Nazanin",              "size": 14, "line": 1.3,  "rtl": True,  "tag": "fa-IR"},
        "ps": {"font": "Bahij Nassim",           "size": 14, "line": 1.3,  "rtl": True,  "tag": "ps-AF"},
        "sd": {"font": "MB Lateefi",             "size": 16, "line": 1.3,  "rtl": True,  "tag": "sd-Arab-PK"},
        "en": {"font": "Times New Roman",        "size": 12, "line": 1.15, "rtl": False, "tag": "en-US"},
        "hi": {"font": "Mangal",                 "size": 13, "line": 1.3,  "rtl": False, "tag": "hi-IN"},
    },
    "default_lang": "ar",
    "body": {"align": "justify", "first_indent_cm": 0.0, "space_before": 0, "space_after": 8},
    "title":    {"scale": 1.6, "bold": True, "space_after": 12},
    "subtitle": {"scale": 1.25, "bold": True, "space_after": 10},
    "heading1": {"scale": 1.3, "bold": True, "space_before": 12, "space_after": 6, "align": "center"},
    "heading2": {"scale": 1.15, "bold": True, "space_before": 10, "space_after": 4, "align": "start"},
    "basmala":  {"scale": 1.25, "bold": True, "space_after": 12},
    "quote":    {"scale": 1.0, "indent_cm": 1.0, "italic": False, "bold": False},   # block quotes / ayat / ahadith
    "verse":    {"scale": 1.0, "gap_cm": 1.0},                                       # poetry (two hemistiches)
    "footnote": {"scale": 0.75, "line": 1.0},
    "lemma":    {"bold": True, "color": None},          # **...**  e.g. "قوله ..." / matn quoted in a hashiya
    "ayah":     {"bold": False, "color": None},         # text inside ﴿ ... ﴾ gets this
    "ref_color": "7F7F7F", "ref_scale": 0.75,           # [[ref:...]] marginal references
    "marker_color": "808080", "marker_scale": 0.75,     # @page original page numbers
    "uncertain_highlight": True,                         # yellow highlight + comment on {{...|...}}
    "page_marker_format": "[ص {n}]",                    # @page argument is the printed digits only
    "comment_author": "Composer",
}
# If a client .docx/.dotx template is supplied, map our style roles to THEIR style names here
# (leave empty to use our own styles).  e.g. {"body": "Matn", "heading1": "Unwan", "footnote": "Footnote Text"}
TEMPLATE_STYLE_MAP = {}
# =====================================================================

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class Composer:
    def __init__(self, template=None):
        self.doc = Document(template) if template else Document()
        self.templ = bool(template)
        if not template:   # clear default empty paragraph
            for p in list(self.doc.paragraphs):
                p._element.getparent().remove(p._element)
        self.lang = STYLE["default_lang"]
        self.fn_part = None
        self.fn_next = 1
        self._table = None
        self._table_cols = None
        self._setup_page()
        self._setup_normal()

    # ---------------- low-level XML helpers ----------------
    def L(self, lang=None):
        return STYLE["langs"][lang or self.lang]

    @staticmethod
    def _rpr_fonts(rpr, font, size, rtl, tag, bold=False, italic=False, color=None):
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), font)
        for t in ("w:b", "w:bCs"):
            if bold and rpr.find(qn(t)) is None: rpr.append(OxmlElement(t))
        for t in ("w:i", "w:iCs"):
            if italic and rpr.find(qn(t)) is None: rpr.append(OxmlElement(t))
        if color:
            c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
        for t in ("w:sz", "w:szCs"):
            e = rpr.find(qn(t))
            if e is None: e = OxmlElement(t); rpr.append(e)
            e.set(qn("w:val"), str(int(round(size * 2))))
        if rtl and rpr.find(qn("w:rtl")) is None:
            rpr.append(OxmlElement("w:rtl"))
        lg = rpr.find(qn("w:lang"))
        if lg is None: lg = OxmlElement("w:lang"); rpr.append(lg)
        lg.set(qn("w:bidi" if rtl else "w:val"), tag)

    @staticmethod
    def _jc(rtl, align):
        # In bidi paragraphs "start" = right. Word accepts left/right/center/both; for RTL right==start.
        m = {"justify": "both", "center": "center",
             "start": "right" if rtl else "left", "end": "left" if rtl else "right"}
        return m.get(align, align)

    def _setup_page(self):
        P = STYLE["page"]
        for s in self.doc.sections:
            s.page_width, s.page_height = Cm(P["width_cm"]), Cm(P["height_cm"])
            s.top_margin, s.bottom_margin = Cm(P["top_cm"]), Cm(P["bottom_cm"])
            s.left_margin, s.right_margin = Cm(P["left_cm"]), Cm(P["right_cm"])
            s.gutter = Cm(P["gutter_cm"])
            s.header_distance, s.footer_distance = Cm(P["header_cm"]), Cm(P["footer_cm"])
            if self.L()["rtl"] and s._sectPr.find(qn("w:bidi")) is None:
                s._sectPr.append(OxmlElement("w:bidi"))
        if P["mirror_margins"]:
            st = self.doc.settings.element
            if st.find(qn("w:mirrorMargins")) is None:
                st.append(OxmlElement("w:mirrorMargins"))

    def _setup_normal(self):
        if self.templ: return
        L = self.L()
        n = self.doc.styles["Normal"]
        self._rpr_fonts(n.element.get_or_add_rPr(), L["font"], L["size"], L["rtl"], L["tag"])
        if L["rtl"]: n.element.get_or_add_pPr().append(OxmlElement("w:bidi"))

    # ---------------- styles (created lazily per language) ----------------
    def pstyle(self, role, lang=None):
        lang = lang or self.lang
        if self.templ and role in TEMPLATE_STYLE_MAP:
            return TEMPLATE_STYLE_MAP[role]
        name = f"{role.capitalize()} ({lang})"
        if name in [s.name for s in self.doc.styles]:
            return name
        L = STYLE["langs"][lang]; rtl = L["rtl"]
        B = STYLE["body"]
        cfg = {"scale": 1.0, "bold": False, "italic": False, "align": B["align"],
               "space_before": B["space_before"], "space_after": B["space_after"],
               "line": L["line"], "first": B["first_indent_cm"], "indent": 0, "color": None, "keep": False}
        if role in ("title", "subtitle", "basmala", "center", "publisher", "ornament", "marker"):
            cfg.update(align="center", first=0)
        if role in STYLE and isinstance(STYLE[role], dict):
            r = STYLE[role]
            cfg.update({k: v for k, v in r.items() if k in cfg})
            if "indent_cm" in r: cfg["indent"] = r["indent_cm"]
        if role in ("title", "subtitle", "heading1", "heading2", "basmala"):
            cfg.update(first=0, keep=True, line=1.0)
        if role == "cont": cfg["first"] = 0
        if role == "quote": cfg["first"] = 0
        if role == "publisher": cfg.update(scale=1.3, bold=True, space_after=0, line=1.1)
        if role == "center": cfg.update(space_after=2)
        if role == "marker":
            cfg.update(scale=STYLE["marker_scale"], color=STYLE["marker_color"], keep=True,
                       space_before=4, space_after=4)
        if role == "footnote":
            cfg.update(scale=STYLE["footnote"]["scale"], line=STYLE["footnote"]["line"],
                       first=0, space_after=0, space_before=0)
        st = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = self.doc.styles["Normal"]; st.quick_style = True
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(cfg["space_before"]), Pt(cfg["space_after"])
        pf.line_spacing = cfg["line"]; pf.keep_with_next = cfg["keep"]
        if cfg["first"]: pf.first_line_indent = Cm(cfg["first"])
        if cfg["indent"]:
            pf.left_indent = Cm(cfg["indent"]); pf.right_indent = Cm(cfg["indent"])
        ppr = st.element.get_or_add_pPr()
        bd = OxmlElement("w:bidi")
        if not rtl: bd.set(qn("w:val"), "0")      # explicit LTR inside an RTL document
        ppr.append(bd)
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), self._jc(rtl, cfg["align"])); ppr.append(jc)
        if role in ("heading1", "heading2"):
            ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), "0" if role == "heading1" else "1"); ppr.append(ol)
        self._rpr_fonts(st.element.get_or_add_rPr(), L["font"], L["size"] * cfg["scale"], rtl, L["tag"],
                        cfg["bold"], cfg["italic"], cfg["color"])
        return name

    def cstyle(self, role):
        name = {"lemma": "Lemma", "ayah": "Ayah", "ref": "Marginal Ref", "blank": "Scribal Blank",
                "fnref": "Footnote Reference"}[role]
        if name in [s.name for s in self.doc.styles]:
            return name
        st = self.doc.styles.add_style(name, WD_STYLE_TYPE.CHARACTER); st.quick_style = True
        rpr = st.element.get_or_add_rPr()
        def add(tag, val=None):
            e = OxmlElement(tag)
            if val is not None: e.set(qn("w:val"), val)
            rpr.append(e)
        if role in ("lemma", "ayah"):
            c = STYLE[role]
            if c.get("bold"): add("w:b"); add("w:bCs")
            if c.get("color"): add("w:color", c["color"])
        elif role in ("ref", "blank"):
            add("w:color", STYLE["ref_color"])
            sz = str(int(self.L()["size"] * STYLE["ref_scale"] * 2)); add("w:sz", sz); add("w:szCs", sz)
        elif role == "fnref":
            add("w:vertAlign", "superscript")
        return name

    # ---------------- footnotes ----------------
    def _ensure_footnotes(self):
        if self.fn_part is not None: return
        dp = self.doc.part
        for rel in dp.rels.values():
            if rel.reltype == RT.FOOTNOTES:
                self.fn_part = rel.target_part
                ids = [int(f.get(qn("w:id"))) for f in self.fn_part.element.findall(qn("w:footnote"))]
                self.fn_next = max(ids + [0]) + 1
                return
        xml = ('<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
               'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
               '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:separator/></w:r></w:p></w:footnote>'
               '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote></w:footnotes>')
        part = XmlPart(PackURI("/word/footnotes.xml"),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                       parse_xml(xml), dp.package)
        dp.relate_to(part, RT.FOOTNOTES)
        self.fn_part = part
        # register separators in settings.xml (Word expects this); must sit before <w:compat>
        st = self.doc.settings.element
        if st.find(qn("w:footnotePr")) is None:
            fp = parse_xml('<w:footnotePr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                           '<w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>')
            anchor = None
            for tag in ("w:endnotePr", "w:compat", "w:docVars", "w:rsids"):
                anchor = st.find(qn(tag))
                if anchor is not None: break
            if anchor is not None: anchor.addprevious(fp)
            else: st.append(fp)

    def add_footnote(self, paragraph, text):
        self._ensure_footnotes()
        fid = self.fn_next; self.fn_next += 1
        L = self.L()
        # reference in body
        r = paragraph.add_run(); self._rtl_run(r)
        r.style = self.cstyle("fnref")
        ref = OxmlElement("w:footnoteReference"); ref.set(qn("w:id"), str(fid)); r._r.append(ref)
        # footnote body
        fn = OxmlElement("w:footnote"); fn.set(qn("w:id"), str(fid))
        p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr")
        ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), self.doc.styles[self.pstyle("footnote")].style_id)
        ppr.append(ps); p.append(ppr)
        r1 = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
        rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), self.doc.styles[self.cstyle("fnref")].style_id); rp.append(rs)
        if L["rtl"]: rp.append(OxmlElement("w:rtl"))
        r1.append(rp); r1.append(OxmlElement("w:footnoteRef")); p.append(r1)
        r2 = OxmlElement("w:r"); rp2 = OxmlElement("w:rPr")
        if L["rtl"]: rp2.append(OxmlElement("w:rtl"))
        r2.append(rp2); t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = " " + text
        r2.append(t); p.append(r2); fn.append(p)
        self.fn_part.element.append(fn)

    # ---------------- runs / inline markup ----------------
    def _rtl_run(self, run):
        if self.L()["rtl"]:
            rpr = run._r.get_or_add_rPr()
            if rpr.find(qn("w:rtl")) is None: rpr.append(OxmlElement("w:rtl"))
        return run

    TOKEN = re.compile(r"(\*\*.+?\*\*|\{\{.+?\}\}|\[\[.+?\]\]|\^\[[\s\S]+?\]\^|﴿.+?﴾)")

    def inline(self, p, text, base_char_style=None):
        for part in self.TOKEN.split(text):
            if not part: continue
            if part.startswith("**") and part.endswith("**"):
                self.inline(p, part[2:-2], self.cstyle("lemma"))
            elif part.startswith("{{") and part.endswith("}}"):
                body = part[2:-2]
                word, _, note = body.partition("|")
                r = self._rtl_run(p.add_run(word, style=base_char_style))
                if STYLE["uncertain_highlight"]:
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                self.doc.add_comment(r, text=note.strip() or "Reading uncertain — please verify.",
                                     author=STYLE["comment_author"], initials="C")
            elif part.startswith("[[") and part.endswith("]]"):
                body = part[2:-2]
                if body.startswith("ref:"):
                    self._rtl_run(p.add_run("[" + body[4:].strip() + "]", style=self.cstyle("ref")))
                    self._rtl_run(p.add_run(" "))
                elif body.startswith("blank"):
                    label = body.partition(":")[2].strip() or ("(بياض)" if self.L()["rtl"] else "(blank)")
                    r = self._rtl_run(p.add_run(label, style=self.cstyle("blank")))
                    self.doc.add_comment(r, text="Gap left blank in the original manuscript.",
                                         author=STYLE["comment_author"], initials="C")
                else:
                    self._rtl_run(p.add_run(part, style=base_char_style))
            elif part.startswith("^[") and part.endswith("]^"):
                self.add_footnote(p, part[2:-2])
            elif part.startswith("﴿"):
                self._rtl_run(p.add_run(part, style=self.cstyle("ayah")))
            else:
                self._rtl_run(p.add_run(part, style=base_char_style))

    def para(self, role, text, page_break=False):
        p = self.doc.add_paragraph(style=self.pstyle(role))
        if page_break: p.paragraph_format.page_break_before = True
        self.inline(p, text)
        return p

    def close_table(self):
        self._table = None
        self._table_cols = None

    def add_row(self, cells):
        """One printed table row. Cells are separated by the caller. Never reorder."""
        L = self.L()
        if self._table is None or len(cells) != self._table_cols:
            self._table = self.doc.add_table(rows=0, cols=len(cells))
            self._table_cols = len(cells)
            tblPr = self._table._tbl.tblPr
            if L["rtl"] and tblPr.find(qn("w:bidiVisual")) is None:
                tblPr.append(OxmlElement("w:bidiVisual"))
        row = self._table.add_row()
        for cell, txt in zip(row.cells, cells):
            cp = cell.paragraphs[0]
            cp.style = self.pstyle("body")
            self.inline(cp, txt)

    def verse(self, a, b):
        L = self.L()
        t = self.doc.add_table(rows=1, cols=2)
        tblPr = t._tbl.tblPr
        if L["rtl"]: tblPr.append(OxmlElement("w:bidiVisual"))
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblPr.append(jc)
        for cell, txt in zip(t.rows[0].cells, (a, b)):
            cp = cell.paragraphs[0]; cp.style = self.pstyle("center")
            self.inline(cp, txt.strip())

    # ---------------- header / footer ----------------
    def _field_run(self, p, instr):
        r = p.add_run()
        for kind in ("begin", None, "end"):
            if kind:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
        return r

    def header_footer(self):
        s = self.doc.sections[0]
        PN = STYLE["page_number"]
        rtl = self.L()["rtl"]
        if STYLE["header_text"]:
            hp = s.header.paragraphs[0]; hp.style = self.pstyle("center")
            self._rtl_run(hp.add_run(STYLE["header_text"]))
        if PN["show"]:
            hf = s.footer if PN["position"] == "footer" else s.header
            p = hf.paragraphs[0] if not (PN["position"] == "header" and STYLE["header_text"]) else hf.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            if rtl: ppr.append(OxmlElement("w:bidi"))
            jc = OxmlElement("w:jc"); jc.set(qn("w:val"), self._jc(rtl, PN["align"])); ppr.append(jc)
            r = self._field_run(p, "PAGE"); r.font.size = Pt(PN["size"])

    # ---------------- transcript parser ----------------
    def build(self, transcript_path):
        lines = Path(transcript_path).read_text(encoding="utf-8").splitlines()
        pending_break = False
        for raw in lines:
            line = raw.strip()
            if not line or line.startswith("#"):       # blank or comment line
                continue
            if line.startswith("@"):
                cmd, _, arg = line[1:].partition(" ")
                cmd = cmd.lower(); arg = arg.strip()
                if cmd not in ("row", "pdf", "skip"):
                    self.close_table()
                if cmd == "lang":
                    if arg not in STYLE["langs"]:
                        raise SystemExit(f"Unknown language '{arg}'. Add it to STYLE['langs'].")
                    self.lang = arg
                elif cmd in ("pdf", "skip"):
                    # coverage ledger only — never printed, never a reason to drop the page's text
                    continue
                elif cmd == "newpage":
                    pending_break = True
                elif cmd == "page":
                    self.para("marker", STYLE["page_marker_format"].format(n=arg), pending_break)
                    pending_break = False
                elif cmd in ("title", "subtitle", "heading1", "heading2", "basmala", "center",
                             "publisher", "quote", "cont"):
                    self.para(cmd, arg, pending_break); pending_break = False
                elif cmd == "ornament":
                    self.para("ornament", arg or "٭ ٭ ٭", pending_break); pending_break = False
                elif cmd == "sheet":
                    parts = arg.lower().split()
                    if len(parts) != 3 or parts[1] != "of" or not parts[0].isdigit() or not parts[2].isdigit():
                        raise SystemExit("@sheet needs '7 of 9'")
                    self.para("marker", f"[شیٹ {parts[0]}/{parts[2]}]", pending_break)
                    pending_break = False
                elif cmd == "verse":
                    a, _, b = arg.partition("|||"); self.verse(a, b)
                elif cmd == "row":
                    cells = [c.strip() for c in arg.split("|||")]
                    if not any(cells):
                        raise SystemExit("@row with no cells")
                    self.add_row(cells)
                else:
                    raise SystemExit(f"Unknown directive @{cmd}")
                continue
            self.close_table()
            self.para("body", line, pending_break); pending_break = False
        self.header_footer()

    # ---------------- schema order fix (Word is strict) ----------------
    def finalize(self, out):
        PPR = ["pStyle","keepNext","keepLines","pageBreakBefore","framePr","widowControl","numPr","suppressLineNumbers",
               "pBdr","shd","tabs","suppressAutoHyphens","kinsoku","wordWrap","overflowPunct","topLinePunct",
               "autoSpaceDE","autoSpaceDN","bidi","adjustRightInd","snapToGrid","spacing","ind","contextualSpacing",
               "mirrorIndents","suppressOverlap","jc","textDirection","textAlignment","textboxTightWrap","outlineLvl",
               "divId","cnfStyle","rPr","sectPr","pPrChange"]
        RPR = ["rStyle","rFonts","b","bCs","i","iCs","caps","smallCaps","strike","dstrike","outline","shadow","emboss",
               "imprint","noProof","snapToGrid","vanish","webHidden","color","spacing","w","kern","position","sz","szCs",
               "highlight","u","effect","bdr","shd","fitText","vertAlign","rtl","cs","em","lang","eastAsianLayout",
               "specVanish","oMath"]
        SECT = ["headerReference","footerReference","footnotePr","endnotePr","type","pgSz","pgMar","paperSrc",
                "pgBorders","lnNumType","pgNumType","cols","formProt","vAlign","noEndnote","titlePg","textDirection",
                "bidi","rtlGutter","docGrid","printerSettings","sectPrChange"]
        TBL = ["tblStyle","tblpPr","tblOverlap","bidiVisual","tblStyleRowBandSize","tblStyleColBandSize","tblW","jc",
               "tblCellSpacing","tblInd","tblBorders","shd","tblLayout","tblCellMar","tblLook","tblCaption",
               "tblDescription","tblPrChange"]
        def srt(el, order):
            kids = list(el)
            kids.sort(key=lambda c: order.index(c.tag.replace(W, "")) if c.tag.replace(W, "") in order else len(order))
            for c in kids: el.remove(c)
            for c in kids: el.append(c)
        def fix(root):
            for el in root.iter(W + "pPr"): srt(el, PPR)
            for el in root.iter(W + "rPr"): srt(el, RPR)
            for el in root.iter(W + "sectPr"): srt(el, SECT)
            for el in root.iter(W + "tblPr"): srt(el, TBL)
        fix(self.doc.styles.element); fix(self.doc.element)
        for s in self.doc.sections:
            fix(s.footer._element); fix(s.header._element)
        if self.fn_part is not None: fix(self.fn_part.element)
        self.doc.save(out)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    c = Composer(sys.argv[3] if len(sys.argv) > 3 else None)
    c.build(sys.argv[1])
    c.finalize(sys.argv[2])
    print("Saved:", sys.argv[2])
